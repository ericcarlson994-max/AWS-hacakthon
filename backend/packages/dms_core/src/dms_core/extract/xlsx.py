from __future__ import annotations

import io
from typing import Any

import openpyxl

from dms_core.models import Block


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _table_text(values: list[str], header: list[str]) -> str:
    pairs: list[str] = []
    for index, value in enumerate(values):
        if not value:
            continue
        column = header[index] if index < len(header) and header[index] else f"col{index + 1}"
        pairs.append(f"{column}: {value}")
    return " | ".join(pairs)


def _sheet_blocks(sheet: Any, sheet_name: str, page: int, order: int) -> tuple[list[Block], int]:
    rows = [values for values in ([_cell_text(value) for value in row] for row in sheet.iter_rows(values_only=True)) if any(values)]
    lead: list[str] = []
    for values in rows:
        filled = [value for value in values if value]
        if len(filled) != 1:
            break
        lead.append(filled[0])
    if len(lead) == len(rows) and len(rows) > 1:
        lead = []
    blocks: list[Block] = []

    def add(block_type: str, text: str, level: int | None = None) -> None:
        nonlocal order
        blocks.append(Block(type=block_type, text=text, page=page, level=level, order=order, source="xlsx"))
        order += 1

    if lead:
        add("heading", lead[0], 1)
        if sheet_name.strip() and sheet_name.strip() != lead[0]:
            add("heading", sheet_name, 2)
        for text in lead[1:]:
            add("paragraph", text)
    else:
        add("heading", sheet_name, 1)
    header: list[str] | None = None
    for values in rows[len(lead):]:
        if header is None:
            header = values
            continue
        add("table", _table_text(values, header))
    return blocks, order


def extract_xlsx(data: bytes, start_order: int = 0) -> tuple[list[Block], int, str | None]:
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    blocks: list[Block] = []
    order = start_order
    try:
        sheet_names = list(workbook.sheetnames)
        for sheet_index, sheet_name in enumerate(sheet_names):
            sheet_blocks, order = _sheet_blocks(workbook[sheet_name], sheet_name, sheet_index + 1, order)
            blocks.extend(sheet_blocks)
        properties = workbook.properties
        raw_title = properties.title if properties is not None else None
        title = (raw_title or "").strip() or None
    finally:
        workbook.close()
    return blocks, len(sheet_names), title
