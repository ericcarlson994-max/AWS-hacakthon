from __future__ import annotations

import io
from uuid import UUID

import pytest
from docx import Document as DocxDocument

from dms_adapters.container import build_fake_container
from dms_core.dedup.exact import sha256_hex
from dms_core.ports import Container
from dms_workers import run as worker_run
from dms_workers.indexing import reindex_document

TOPICS = [
    "procurement thresholds",
    "vendor registration",
    "tender evaluation",
    "contract signing",
    "payment schedules",
    "audit requirements",
    "conflict of interest",
    "record keeping",
    "emergency purchases",
    "appeals process",
]


def section_body(topic: str, variant: str = "") -> str:
    sentences = [
        f"This section on {topic} sets out the rules every officer must follow {variant}.",
        f"Officers handling {topic} shall document each decision with supporting evidence and approvals.",
        f"Any deviation regarding {topic} must be reported to the head of department within fourteen days.",
        f"Training on {topic} is mandatory for new staff and refreshed every two years for existing staff.",
    ]
    return " ".join(sentences * 3)


def build_docx(title: str, sections: list[tuple[str, str]], preface: list[str] | None = None) -> bytes:
    document = DocxDocument()
    document.add_heading(title, level=1)
    for line in preface or []:
        document.add_paragraph(line)
    for heading, body in sections:
        document.add_heading(heading, level=2)
        document.add_paragraph(body)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def ten_section_docx(changed_index: int | None = None) -> bytes:
    sections = []
    for position, topic in enumerate(TOPICS):
        variant = "with the amendments approved in 2025" if position == changed_index else ""
        sections.append((f"{position + 1}. {topic.title()}", section_body(topic, variant)))
    return build_docx("Procurement Policy", sections)


@pytest.fixture
def container() -> Container:
    c = build_fake_container()
    c.repo.upsert_department("finance", "Finance")
    c.repo.upsert_user("u1", "alice", "Alice")
    c.repo.set_membership("u1", "finance", "contributor")
    return c


def upload(c: Container, data: bytes, filename: str, document_id: UUID | None = None) -> tuple[UUID, UUID]:
    if document_id is None:
        document_id = c.repo.create_document(department_id="finance", created_by="u1").id
    sha = sha256_hex(data)
    c.blobs.put(sha, data, "application/octet-stream")
    version = c.repo.create_version(
        document_id=document_id,
        blob_sha256=sha,
        original_filename=filename,
        mime_type="application/octet-stream",
        size_bytes=len(data),
        uploaded_by="u1",
    )
    c.jobs.enqueue(document_id, version.id, "extract")
    return document_id, version.id


def test_first_version_runs_full_pipeline_to_ready(container: Container) -> None:
    document_id, version_id = upload(container, ten_section_docx(), "procurement_policy.docx")
    processed = worker_run.drain(container)
    assert processed == 4
    doc = container.repo.get_document(document_id)
    assert doc.status == "READY"
    assert doc.root_summary
    assert doc.title == "Procurement Policy"
    assert doc.doctype is not None
    version = container.repo.get_version(version_id)
    assert version.chunk_manifest and len(version.chunk_manifest) >= 10
    assert version.merkle_root
    chunks = container.repo.get_chunks(version.chunk_manifest)
    assert all(chunk.embedding for chunk in chunks)
    nodes = container.repo.get_raptor_nodes(document_id, 1)
    assert nodes
    index = container.index
    assert str(document_id) in index.titles
    assert index.titles[str(document_id)]["filename"] == "procurement_policy.docx"
    indexed = [d for d in index.chunks.values() if d["document_id"] == str(document_id)]
    leaf_docs = [d for d in indexed if d["level"] == 0]
    summary_docs = [d for d in indexed if d["level"] >= 1]
    assert len(leaf_docs) == len(version.chunk_manifest)
    assert len(summary_docs) == len([n for n in nodes if n.level >= 1])
    assert all(d["allowed_groups"] == ["finance"] for d in indexed)
    assert all(d["embedding"] for d in indexed)
    assert reindex_document(container, document_id) == len(indexed)


def test_new_version_reuses_unchanged_chunks(container: Container) -> None:
    document_id, _ = upload(container, ten_section_docx(), "procurement_policy.docx")
    worker_run.drain(container)
    first_root = container.repo.get_document(document_id).root_summary
    embed_calls_before = sum(len(call) for call in container.embedder.calls)
    _, second_version_id = upload(container, ten_section_docx(changed_index=3), "procurement_policy_v2.docx", document_id)
    worker_run.drain(container)
    doc = container.repo.get_document(document_id)
    assert doc.status == "READY"
    assert doc.root_summary and first_root
    version = container.repo.get_version(second_version_id)
    stats = version.delta_stats
    assert stats is not None
    assert stats.added >= 1
    assert stats.reused / (stats.reused + stats.added) >= 0.8
    assert stats.reembedded == stats.added
    assert stats.raptor_nodes_refreshed >= 1
    assert container.repo.get_raptor_nodes(document_id, 2)
    previous = container.repo.get_version_by_no(document_id, 1)
    assert set(previous.chunk_manifest) - set(version.chunk_manifest)
    assert container.repo.get_chunks(previous.chunk_manifest)
    indexed_ids = {d["chunk_id"] for d in container.index.chunks.values() if d["level"] == 0}
    assert set(version.chunk_manifest) <= indexed_ids
    assert not (set(previous.chunk_manifest) - set(version.chunk_manifest)) & indexed_ids
    assert sum(len(call) for call in container.embedder.calls) > embed_calls_before
    stages = [job.stage for job in container.repo.list_jobs(document_id)]
    assert stages.count("delta") == 1


def test_circular_supersedes_older_document(container: Container) -> None:
    old_bytes = build_docx(
        "Circular on Travel Claims",
        [("Rules", section_body("travel claims"))],
        preface=["Reference: JPA/CIRC/2023/01", "Effective 2023-01-01"],
    )
    new_bytes = build_docx(
        "Circular on Travel Claims Revised",
        [("Rules", section_body("revised travel claims", "from now on"))],
        preface=["Reference: JPA/CIRC/2024/05", "This circular supersedes JPA/CIRC/2023/01", "Effective 2024-06-01"],
    )
    old_id, _ = upload(container, old_bytes, "circular_2023.docx")
    worker_run.drain(container)
    new_id, _ = upload(container, new_bytes, "circular_2024.docx")
    worker_run.drain(container)
    old_doc = container.repo.get_document(old_id)
    new_doc = container.repo.get_document(new_id)
    assert old_doc.meta.reference_no == "JPA/CIRC/2023/01"
    assert new_doc.meta.supersedes_ref == "JPA/CIRC/2023/01"
    assert new_doc.meta.effective_date == "2024-06-01"
    assert old_doc.superseded_by == new_id
    assert new_doc.superseded_by is None
    assert container.index.titles[str(old_id)]["superseded"] is True
    old_chunks = [d for d in container.index.chunks.values() if d["document_id"] == str(old_id)]
    assert old_chunks and all(d["superseded"] for d in old_chunks)


def test_near_duplicate_is_flagged_but_still_processed(container: Container) -> None:
    sections = [(f"Part {topic}", section_body(topic)) for topic in TOPICS[:4]]
    original_id, _ = upload(container, build_docx("Records Guideline", sections), "records.docx")
    worker_run.drain(container)
    copy_id, _ = upload(container, build_docx("Records Guideline Copy", sections), "records_copy.docx")
    worker_run.drain(container)
    copy = container.repo.get_document(copy_id)
    assert copy.status == "READY"
    assert copy.duplicate_tier == "near"
    assert copy.duplicate_of == original_id


def test_ai_tags_are_suggested_from_taxonomy(container: Container) -> None:
    tag = container.repo.get_or_create_tag("procurement", "finance")
    container.repo.set_tag_centroid(tag.id, container.embedder.embed([section_body("procurement thresholds")])[0])
    document_id, _ = upload(container, build_docx("Thresholds", [("Rules", section_body("procurement thresholds"))]), "t.docx")
    worker_run.drain(container)
    refs = container.repo.get_document_tags(document_id)
    assert [(ref.name, ref.source, ref.status) for ref in refs] == [("procurement", "ai", "suggested")]
    assert container.index.titles[str(document_id)]["tags"] == ["procurement"]


def test_failing_stage_retries_then_marks_failed(container: Container, monkeypatch: pytest.MonkeyPatch) -> None:
    attempts: list[int] = []

    def broken_extract(c: Container, job: object) -> None:
        attempts.append(1)
        raise RuntimeError("blob unreadable")

    monkeypatch.setitem(worker_run.STAGES, "extract", broken_extract)
    document_id, _ = upload(container, ten_section_docx(), "broken.docx")
    worker_run.drain(container)
    assert len(attempts) == container.settings.job_max_attempts
    doc = container.repo.get_document(document_id)
    assert doc.status == "FAILED"
    assert "blob unreadable" in (doc.status_detail or "")
    jobs = container.repo.list_jobs(document_id)
    assert [job.status for job in jobs] == ["failed"]


def test_unsupported_file_marks_document_failed(container: Container) -> None:
    document_id, _ = upload(container, b"just some bytes", "notes.bin")
    worker_run.drain(container)
    assert container.repo.get_document(document_id).status == "FAILED"


def test_run_once_returns_false_on_empty_queue(container: Container) -> None:
    assert worker_run.run_once(container) is False
    assert worker_run.drain(container) == 0


def auto_filed_folder(c: Container, data: bytes, filename: str, department_id: str = "finance", folder_id: str | None = None) -> str | None:
    document_id = c.repo.create_document(department_id=department_id, created_by="u1", folder_id=folder_id).id
    upload(c, data, filename, document_id)
    worker_run.drain(c)
    return c.repo.get_document(document_id).folder_id


def simple_docx(title: str, preface: list[str] | None = None) -> bytes:
    return build_docx(title, [("Scope", section_body(title.lower()))], preface=preface)


def test_auto_filing_by_doctype(container: Container) -> None:
    container.repo.upsert_department("hr", "HR")
    assert auto_filed_folder(container, simple_docx("Procurement Policy"), "policy_procurement.docx") == "policies"
    assert auto_filed_folder(container, simple_docx("Claims SOP"), "sop_claims.docx", folder_id="inbox") == "sop-fin"
    assert auto_filed_folder(container, simple_docx("Leave SOP"), "sop_leave.docx", department_id="hr") == "sop-hr"
    circular_2024 = simple_docx("Remote Work Circular", ["Effective 2024-07-01"])
    assert auto_filed_folder(container, circular_2024, "circular_remote.docx") == "circ-2024"
    circular_2025 = simple_docx("Travel Circular", ["Effective 2025-02-01"])
    assert auto_filed_folder(container, circular_2025, "circular_travel.docx") == "circulars"
    assert auto_filed_folder(container, simple_docx("Board Minutes"), "minutes_board.docx") == "min-mgmt"
    assert auto_filed_folder(container, simple_docx("Annual Report"), "report_annual.docx") == "reports"
    assert auto_filed_folder(container, simple_docx("Random Notes"), "notes.docx") == "inbox"


def test_auto_filing_keeps_user_chosen_folder(container: Container) -> None:
    assert auto_filed_folder(container, simple_docx("Procurement Policy"), "policy_x.docx", folder_id="reports") == "reports"


def test_auto_filing_falls_back_when_subfolder_deleted(container: Container) -> None:
    container.repo.delete_folder("sop-fin")
    assert auto_filed_folder(container, simple_docx("Claims SOP"), "sop_claims.docx") == "sops"
