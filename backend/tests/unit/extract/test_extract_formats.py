from __future__ import annotations

import io

import docx as python_docx
import openpyxl
import pytest
from PIL import Image
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from dms_adapters.fakes.ai import FakeVisionOcr
from dms_core.extract.filetype import (
    DOCX_MIME,
    PDF_MIME,
    SUPPORTED_MIME,
    XLSX_MIME,
    UnsupportedFileType,
    detect_mime,
)
from dms_core.extract.markdown import parse_ocr_markdown
from dms_core.extract.router import extract

BODY_LINES = [
    "This policy sets out the procedure for handling official correspondence in the agency.",
    "All officers must register incoming letters within two working days of receipt.",
    "Records shall be retained for a minimum of seven years unless otherwise directed.",
]


def make_pdf(pages: list[list[tuple[str, int]] | None], title: str | None = None) -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    if title:
        pdf.setTitle(title)
    for page_lines in pages:
        if page_lines is None:
            pdf.rect(100, 400, 200, 150, fill=1)
        else:
            y = 780
            for text, size in page_lines:
                pdf.setFont("Helvetica", size)
                pdf.drawString(72, y, text)
                y -= size * 3
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def make_png(color: str = "white") -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (40, 30), color).save(buffer, format="PNG")
    return buffer.getvalue()


def make_jpeg() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (40, 30), "gray").save(buffer, format="JPEG")
    return buffer.getvalue()


def make_docx() -> bytes:
    document = python_docx.Document()
    document.core_properties.title = "Correspondence Policy"
    document.add_heading("Correspondence Policy", level=0)
    document.add_heading("Scope", level=1)
    document.add_paragraph("Applies to all officers.")
    document.add_heading("Retention", level=2)
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Record"
    table.cell(0, 1).text = "Years"
    table.cell(1, 0).text = "Letters"
    table.cell(1, 1).text = "7"
    table.cell(2, 0).text = "Memos"
    table.cell(2, 1).text = "3"
    document.add_paragraph("End of document.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_xlsx() -> bytes:
    workbook = openpyxl.Workbook()
    first = workbook.active
    first.title = "Budget"
    first.append(["Item", "Amount"])
    first.append(["Laptops", 1200])
    first.append([None, None])
    first.append(["Chairs", 300])
    second = workbook.create_sheet("Staff")
    second.append(["Name", "Grade"])
    second.append(["Aisyah", "N41"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_pdf_with_blank_page_uses_ocr_once():
    data = make_pdf([[(line, 11) for line in BODY_LINES], None, [(line, 11) for line in BODY_LINES]])
    ocr = FakeVisionOcr()
    result = extract(data, "memo.pdf", ocr)
    assert result.mime_type == PDF_MIME
    assert result.page_count == 3
    assert len(ocr.calls) == 1
    assert result.ocr_pages == [2]
    page_two = [block for block in result.blocks if block.page == 2]
    assert page_two and all(block.source == "ocr" for block in page_two)
    assert page_two[0].type == "heading" and page_two[0].text == "Scanned page"
    orders = [block.order for block in result.blocks]
    assert orders == sorted(orders) and len(set(orders)) == len(orders)
    assert [block.page for block in result.blocks] == sorted(block.page for block in result.blocks)


def test_pdf_blank_page_without_ocr_does_not_crash():
    data = make_pdf([[(line, 11) for line in BODY_LINES], None])
    result = extract(data, "memo.pdf", None)
    assert result.page_count == 2
    assert result.ocr_pages == []
    assert not [block for block in result.blocks if block.page == 2]


def test_pdf_heading_detection_and_metadata_title():
    lines = [("Records Management Policy", 24), ("Section One Overview", 16)] + [(line, 11) for line in BODY_LINES]
    data = make_pdf([lines], title="Records Policy 2026")
    ocr = FakeVisionOcr()
    result = extract(data, "policy.pdf", ocr)
    assert ocr.calls == []
    assert result.embedded_title == "Records Policy 2026"
    headings = [block for block in result.blocks if block.type == "heading"]
    assert [(block.text, block.level) for block in headings] == [
        ("Records Management Policy", 1),
        ("Section One Overview", 2),
    ]
    paragraphs = [block for block in result.blocks if block.type == "paragraph"]
    assert len(paragraphs) == 3
    assert all(block.source == "text_layer" and block.bbox for block in result.blocks)


def test_docx_headings_and_table():
    result = extract(make_docx(), "policy.docx", None)
    assert result.mime_type == DOCX_MIME
    assert result.page_count == 1
    assert result.embedded_title == "Correspondence Policy"
    summary = [(block.type, block.level, block.text) for block in result.blocks]
    assert summary == [
        ("heading", 1, "Correspondence Policy"),
        ("heading", 1, "Scope"),
        ("paragraph", None, "Applies to all officers."),
        ("heading", 2, "Retention"),
        ("table", None, "Record: Letters | Years: 7\nRecord: Memos | Years: 3"),
        ("paragraph", None, "End of document."),
    ]
    assert all(block.source == "docx" for block in result.blocks)
    assert [block.order for block in result.blocks] == list(range(6))


def test_xlsx_rows_serialized():
    result = extract(make_xlsx(), "budget.xlsx", None)
    assert result.mime_type == XLSX_MIME
    assert result.page_count == 2
    summary = [(block.type, block.page, block.text) for block in result.blocks]
    assert summary == [
        ("heading", 1, "Budget"),
        ("table", 1, "Item: Laptops | Amount: 1200"),
        ("table", 1, "Item: Chairs | Amount: 300"),
        ("heading", 2, "Staff"),
        ("table", 2, "Name: Aisyah | Grade: N41"),
    ]
    assert all(block.source == "xlsx" for block in result.blocks)


def test_image_extraction_uses_ocr():
    ocr = FakeVisionOcr(text="# Notice\n\nOffice closed on Friday.")
    result = extract(make_jpeg(), "scan.jpg", ocr)
    assert result.mime_type == "image/jpeg"
    assert result.ocr_pages == [1]
    assert len(ocr.calls) == 1
    assert [(block.type, block.text) for block in result.blocks] == [
        ("heading", "Notice"),
        ("paragraph", "Office closed on Friday."),
    ]


def test_image_without_ocr_returns_empty():
    result = extract(make_png(), "scan.png", None)
    assert result.blocks == []
    assert result.page_count == 1


def test_detect_mime_by_magic_bytes():
    assert detect_mime(make_png(), "fake.pdf") == "image/png"
    assert detect_mime(make_jpeg(), "photo.bin") == "image/jpeg"
    assert detect_mime(make_pdf([[("hello", 11)]]), "noext") == PDF_MIME
    assert detect_mime(make_docx(), "a.xlsx") == DOCX_MIME
    assert detect_mime(make_xlsx(), "a.docx") == XLSX_MIME
    assert SUPPORTED_MIME == {PDF_MIME, DOCX_MIME, XLSX_MIME, "image/png", "image/jpeg"}


def test_detect_mime_rejects_unknown():
    with pytest.raises(UnsupportedFileType):
        detect_mime(b"just some text", "notes.txt")
    with pytest.raises(ValueError):
        detect_mime(b"PK\x03\x04garbage", "archive.zip")


def test_markdown_parser():
    text = "\n".join(
        [
            "# Title",
            "Intro line one",
            "intro line two",
            "",
            "## Section",
            "| Name | Grade |",
            "|---|:---:|",
            "| Ali | N41 |",
            "Closing paragraph.",
        ]
    )
    blocks = parse_ocr_markdown(text, page=3, start_order=10)
    assert [(block.type, block.level, block.text) for block in blocks] == [
        ("heading", 1, "Title"),
        ("paragraph", None, "Intro line one\nintro line two"),
        ("heading", 2, "Section"),
        ("table", None, "Name | Grade\nAli | N41"),
        ("paragraph", None, "Closing paragraph."),
    ]
    assert [block.order for block in blocks] == [10, 11, 12, 13, 14]
    assert all(block.page == 3 and block.source == "ocr" for block in blocks)


def test_xlsx_title_row_becomes_heading_before_sheet_name():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Q3 Summary"
    sheet["A1"] = "Quarterly Budget"
    sheet["A2"] = "Finance Division"
    sheet.append([])
    sheet.append(["Item", "Amount"])
    sheet.append(["Laptops", 1200])
    buffer = io.BytesIO()
    workbook.save(buffer)
    result = extract(buffer.getvalue(), "budget.xlsx", None)
    summary = [(block.type, block.level, block.text) for block in result.blocks]
    assert summary == [
        ("heading", 1, "Quarterly Budget"),
        ("heading", 2, "Q3 Summary"),
        ("paragraph", None, "Finance Division"),
        ("table", None, "Item: Laptops | Amount: 1200"),
    ]
