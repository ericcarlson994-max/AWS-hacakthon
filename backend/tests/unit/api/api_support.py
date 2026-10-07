from __future__ import annotations

import io
from typing import Any

from docx import Document as DocxDocument
from fastapi.testclient import TestClient

from dms_core.ports import Container

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

DEPARTMENTS = {"finance": "Finance", "hr": "Human Resources", "public": "Public"}
USERS: dict[str, tuple[str, list[tuple[str, str]]]] = {
    "alice": ("Alice Tan", [("finance", "contributor"), ("public", "viewer")]),
    "ben": ("Ben Lim", [("hr", "contributor"), ("public", "viewer")]),
    "chloe": ("Chloe Ng", [("finance", "viewer"), ("hr", "viewer"), ("public", "viewer")]),
    "admin": ("Admin", [("finance", "admin"), ("hr", "admin"), ("public", "admin")]),
}


def seed(container: Container) -> None:
    repo = container.repo
    for department_id, name in DEPARTMENTS.items():
        repo.upsert_department(department_id, name)
    for username, (display_name, memberships) in USERS.items():
        repo.upsert_user(username, username, display_name)
        for department_id, role in memberships:
            repo.set_membership(username, department_id, role)


def make_docx(heading: str, paragraphs: list[str]) -> bytes:
    document = DocxDocument()
    document.add_heading(heading, level=1)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class Session:
    def __init__(self, client: TestClient, username: str) -> None:
        self.client = client
        response = client.post("/auth/dev-login", json={"username": username})
        assert response.status_code == 200, response.text
        self.token = response.json()["token"]
        self.headers = {"Authorization": f"Bearer {self.token}"}

    def get(self, url: str, **kwargs: Any) -> Any:
        return self.client.get(url, headers=self.headers, **kwargs)

    def post(self, url: str, **kwargs: Any) -> Any:
        return self.client.post(url, headers=self.headers, **kwargs)

    def put(self, url: str, **kwargs: Any) -> Any:
        return self.client.put(url, headers=self.headers, **kwargs)

    def patch(self, url: str, **kwargs: Any) -> Any:
        return self.client.patch(url, headers=self.headers, **kwargs)

    def delete(self, url: str, **kwargs: Any) -> Any:
        return self.client.delete(url, headers=self.headers, **kwargs)

    def upload(self, department_id: str, filename: str, data: bytes, mime: str = DOCX_MIME) -> Any:
        return self.post("/documents", files={"file": (filename, data, mime)}, data={"department_id": department_id})
