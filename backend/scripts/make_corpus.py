from __future__ import annotations

import json
import shutil
from pathlib import Path

import openpyxl
from docx import Document as DocxDocument
from openpyxl.styles import Font
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas as pdf_canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

AGENCY = "Jabatan Perkhidmatan Digital Negeri (JPDN)"
CORPUS_DIR = Path(__file__).resolve().parent.parent / "samples" / "corpus"

Section = tuple[str, list[str]]

CJK_FONT_CANDIDATES = [
    "/System/Library/Fonts/PingFang.ttc",
    "/System/Library/Fonts/STHeiti Light.ttc",
    "/System/Library/Fonts/Hiragino Sans GB.ttc",
    "/System/Library/Fonts/STHeiti Medium.ttc",
    "/Library/Fonts/Arial Unicode.ttf",
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
]

LATIN_FONT_CANDIDATES = [
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/System/Library/Fonts/Helvetica.ttc",
    "/Library/Fonts/Arial.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
]


def load_font(candidates: list[str], size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    for candidate in candidates:
        if Path(candidate).exists():
            try:
                return ImageFont.truetype(candidate, size)
            except OSError:
                continue
    return ImageFont.load_default(size=size)


def pdf_styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("DocTitle", parent=base["Title"], fontSize=20, leading=24, spaceAfter=10),
        "subtitle": ParagraphStyle("DocSubtitle", parent=base["Normal"], fontSize=10, leading=13, textColor=colors.HexColor("#444444")),
        "h1": ParagraphStyle("H1", parent=base["Heading1"], fontSize=14, leading=18, spaceBefore=10, spaceAfter=6),
        "body": ParagraphStyle("Body", parent=base["BodyText"], fontSize=10.5, leading=14.5, spaceAfter=6),
        "small": ParagraphStyle("Small", parent=base["Normal"], fontSize=8.5, leading=11, textColor=colors.HexColor("#555555")),
    }


def build_pdf(path: Path, title: str, story: list, footer_text: str) -> None:
    def draw_footer(canvas: pdf_canvas.Canvas, document: SimpleDocTemplate) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawString(20 * mm, 12 * mm, footer_text)
        canvas.drawRightString(190 * mm, 12 * mm, f"Page {document.page}")
        canvas.restoreState()

    document = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=20 * mm,
        bottomMargin=22 * mm,
        title=title,
        author=AGENCY,
        invariant=1,
    )
    document.build(story, onFirstPage=draw_footer, onLaterPages=draw_footer)


def sections_story(sections: list[Section], styles: dict[str, ParagraphStyle]) -> list:
    story: list = []
    for heading, paragraphs in sections:
        story.append(Paragraph(heading, styles["h1"]))
        for paragraph in paragraphs:
            story.append(Paragraph(paragraph, styles["body"]))
    return story


def header_block(title: str, lines: list[str], styles: dict[str, ParagraphStyle]) -> list:
    story: list = [Paragraph(AGENCY, styles["subtitle"]), Spacer(1, 4 * mm), Paragraph(title, styles["title"])]
    for line in lines:
        story.append(Paragraph(line, styles["subtitle"]))
    story.append(Spacer(1, 6 * mm))
    return story


def styled_table(rows: list[list[str]], widths: list[float]) -> Table:
    table = Table(rows, colWidths=widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDE6F0")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#888888")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    return table


REMOTE_WORK_2024_SECTIONS: list[Section] = [
    (
        "1. Purpose",
        [
            "This circular sets out the arrangements under which officers of Jabatan Perkhidmatan Digital Negeri may perform their duties from home. It replaces the temporary arrangements that applied during the 2023 office renovation and establishes a single, consistent remote work standard for every division of the department.",
            "The arrangement aims to improve work-life balance, reduce commuting time and support business continuity, while ensuring that service delivery to the public is not affected. Heads of division remain accountable for the output of officers working remotely.",
        ],
    ),
    (
        "2. Effective Date and Scope",
        [
            "This circular takes effect on 1 July 2024 and applies to all permanent, contract and seconded officers of grades 19 to 54 whose duties can be performed without physical presence at the office.",
            "Frontline counter staff, officers on probation during their first three months, and officers assigned to the data centre operations rota are not eligible for remote work under this circular unless approved in writing by the Director.",
        ],
    ),
    (
        "3. Remote Work Entitlement",
        [
            "Eligible officers may work from home for a maximum of <b>two (2) days per week</b>. The remaining working days must be spent at the officer's assigned office. Remote work days may not be accumulated or carried forward to the following week.",
            "Remote work days must be agreed with the immediate supervisor at least one week in advance and recorded in the division's attendance roster. Supervisors must ensure that at least half of each unit is physically present on every working day.",
        ],
    ),
    (
        "4. Working Hours and Availability",
        [
            "Officers working remotely must observe their normal working hours, including the core hours of 9.00 a.m. to 4.00 p.m., and must remain contactable by telephone and the departmental messaging platform throughout those hours.",
            "Officers must attend scheduled meetings by video conference with their camera available on request. Failure to respond to supervisors within thirty minutes during core hours without reasonable cause may result in the remote work day being recorded as leave.",
        ],
    ),
    (
        "5. Equipment and Information Security",
        [
            "Officers must use only department-issued laptops connected through the departmental VPN. Official documents classified as Restricted or above must not be printed, stored or viewed at home.",
            "Officers are responsible for ensuring that their home working area is private and that screens cannot be viewed by other persons. Any loss of equipment or suspected security incident must be reported to the ICT Security Unit within two hours.",
        ],
    ),
    (
        "6. Performance Monitoring",
        [
            "Supervisors will set weekly deliverables for officers working remotely and review progress at the end of each week. Remote work arrangements may be suspended for any officer whose performance falls below the expected standard.",
            "Divisions shall report the number of officers using remote work and any service delivery issues to the Human Resources Division every quarter.",
        ],
    ),
    (
        "7. Enquiries",
        [
            "Enquiries regarding this circular may be directed to the Human Resources Division, Jabatan Perkhidmatan Digital Negeri, at hr.enquiry@jpdn.example.gov.",
        ],
    ),
]


REMOTE_WORK_2026_SECTIONS: list[Section] = [
    (
        "1. Purpose",
        [
            "This circular updates the remote work arrangements for officers of Jabatan Perkhidmatan Digital Negeri following the review of the 2024 arrangements carried out by the Human Resources Division in late 2025.",
            "This circular supersedes PKP/HR/2024/07. All arrangements made under PKP/HR/2024/07 cease to apply from the effective date of this circular.",
        ],
    ),
    (
        "2. Effective Date and Scope",
        [
            "This circular takes effect on 1 March 2026 and applies to all permanent, contract and seconded officers of grades 19 to 54 whose duties can be performed without physical presence at the office.",
            "Officers on probation are now eligible after completing six weeks of service. Frontline counter staff remain ineligible, but may apply for a compressed work week under a separate circular.",
        ],
    ),
    (
        "3. Remote Work Entitlement",
        [
            "Eligible officers may work from home for a maximum of <b>three (3) days per week</b>, an increase from the two days per week allowed under the previous circular. At least two working days per week must be spent at the assigned office.",
            "Each division must designate one common anchor day per week on which all officers of the division are present in the office for planning meetings, team activities and knowledge sharing.",
        ],
    ),
    (
        "4. Working Hours and Availability",
        [
            "Officers working remotely must observe the core hours of 9.30 a.m. to 3.30 p.m. and remain reachable on the departmental messaging platform. Flexible start times between 7.30 a.m. and 9.30 a.m. continue to apply.",
            "Officers must keep their calendar up to date so that colleagues can see their availability. Response time expectations for messages during core hours are thirty minutes.",
        ],
    ),
    (
        "5. Equipment, Allowance and Information Security",
        [
            "Officers working remotely receive a monthly connectivity allowance of RM50, claimable through the standard claims submission procedure. Only department-issued laptops connected through the departmental VPN may be used for official work.",
            "Documents classified as Restricted or above must not be accessed outside the office network. Security incidents must be reported to the ICT Security Unit within two hours.",
        ],
    ),
    (
        "6. Performance and Review",
        [
            "Remote work is a privilege and not an entitlement. Supervisors will review deliverables every fortnight and may reduce or suspend remote work days for officers who do not meet agreed outcomes.",
            "The Human Resources Division will review the effectiveness of this arrangement twelve months after the effective date and report to the Management Committee.",
        ],
    ),
    (
        "7. Enquiries",
        [
            "Enquiries regarding this circular may be directed to the Human Resources Division at hr.enquiry@jpdn.example.gov.",
        ],
    ),
]


PROCUREMENT_SECTIONS_COMMON: list[Section] = [
    (
        "1. Introduction",
        [
            "This policy governs the procurement of goods, services and works by Jabatan Perkhidmatan Digital Negeri. It is issued under the authority of the Director and applies to every division, unit and project office of the department.",
            "The policy ensures that public funds are spent with integrity, transparency and value for money, and that procurement decisions can be audited and justified to the State Treasury and the Audit Department.",
        ],
    ),
    (
        "2. Principles",
        [
            "All procurement activities shall observe the principles of public accountability, transparency, value for money, open and fair competition, and fair dealing. Officers involved in procurement must declare any conflict of interest before participating in an evaluation.",
            "Requirements must not be split into smaller purchases to avoid the thresholds in this policy. Deliberate splitting of requirements is a disciplinary offence and will be referred to the Integrity Unit.",
        ],
    ),
    (
        "3. Roles and Responsibilities",
        [
            "The requesting division prepares the specification, cost estimate and justification. The Finance Division verifies budget availability and manages the procurement process. The Procurement Committee approves purchases above the direct purchase threshold.",
            "The Head of Finance is the controlling officer for procurement records and must ensure that all files are complete and retained for at least seven years.",
        ],
    ),
]

PROCUREMENT_SECTION_4_V1: Section = (
    "4. Procurement Thresholds",
    [
        "Direct purchase may be used for requirements with a total value of up to RM20,000 per item or contract, supported by at least one written quotation.",
        "Quotations from at least three registered suppliers are required for requirements above RM20,000 and up to RM200,000. The quotation exercise must remain open for at least seven days.",
        "Open tender is required for all requirements above RM200,000 and must be advertised on the State e-procurement portal for at least fourteen days.",
    ],
)

PROCUREMENT_SECTION_4_V2: Section = (
    "4. Procurement Thresholds",
    [
        "Direct purchase may be used for requirements with a total value of up to RM50,000 per item or contract, supported by at least two written quotations where available.",
        "Quotations from at least five registered suppliers are required for requirements above RM50,000 and up to RM500,000. The quotation exercise must remain open for at least ten days and must be published on the State e-procurement portal.",
        "Open tender is required for all requirements above RM500,000 and must be advertised for at least twenty-one days. These revised thresholds apply to requisitions raised on or after 1 September 2026.",
    ],
)

PROCUREMENT_SECTIONS_TAIL: list[Section] = [
    (
        "5. Evaluation of Offers",
        [
            "Offers shall be evaluated by a committee of at least three officers who are not involved in preparing the specification. Technical evaluation must be completed before financial evaluation.",
            "The evaluation report must record the scoring of each offer, the reasons for rejecting non-compliant offers, and the recommendation to the approving authority.",
        ],
    ),
    (
        "6. Contract Management",
        [
            "Every contract must name a contract manager responsible for monitoring deliverables, acceptance and payment milestones. Variations above ten percent of the contract value require fresh approval from the Procurement Committee.",
            "Contract managers must record supplier performance at the end of each contract so that the information can be used in future evaluations.",
        ],
    ),
    (
        "7. Emergency Procurement",
        [
            "In an emergency that threatens public safety or the continuity of essential digital services, the Director may approve direct procurement regardless of value. The justification must be documented within three working days.",
            "Emergency procurement must be reported to the Procurement Committee at its next meeting and to the State Treasury within thirty days.",
        ],
    ),
    (
        "8. Records, Audit and Review",
        [
            "All procurement documents, including specifications, quotations, evaluation reports and contracts, must be stored in the document management system with the correct department and tags.",
            "This policy will be reviewed every two years or earlier if the State Treasury issues new instructions. Enquiries may be directed to the Finance Division.",
        ],
    ),
]


CLAIMS_STEPS = [
    "Log in to the staff self-service portal using your departmental account.",
    "Select <i>Claims</i> and choose the claim type: travel, overtime, connectivity allowance or training.",
    "Enter the claim details, including dates, locations, purpose and amount for each item.",
    "Upload scanned receipts in PDF or JPEG format. Each receipt must be legible and show the date and amount.",
    "Submit the claim to your immediate supervisor for verification within 30 days of the expense.",
    "The supervisor verifies the claim and forwards it to the Finance Division within five working days.",
    "The Finance Division checks the claim against the rates in the table below and approves or returns it with reasons.",
    "Approved claims are paid together with the next monthly salary. Claims approved after the 15th of the month are paid in the following month.",
]

CLAIMS_RATES = [
    ["Claim type", "Rate", "Supporting documents", "Approver"],
    ["Mileage (car)", "RM0.60 per km", "Travel log, map printout", "Supervisor"],
    ["Hotel lodging", "Up to RM250 per night", "Hotel invoice", "Head of Division"],
    ["Meal allowance", "RM30 per day", "Not required", "Supervisor"],
    ["Overtime", "Per salary schedule", "Approved overtime form", "Head of Division"],
    ["Connectivity allowance", "RM50 per month", "Remote work roster", "Supervisor"],
    ["Training fees", "Actual cost", "Invoice, attendance letter", "Finance Division"],
]


def claims_sections() -> tuple[list[Section], list[Section]]:
    before = [
        (
            "1. Purpose",
            [
                "This standard operating procedure describes how officers of Jabatan Perkhidmatan Digital Negeri submit claims for official expenses, and how those claims are verified, approved and paid by the Finance Division.",
                "Following this procedure ensures that claims are paid accurately and on time, and that every payment can be supported during an internal or external audit.",
            ],
        ),
        (
            "2. Scope",
            [
                "This SOP applies to travel claims, overtime claims, the remote work connectivity allowance and training fee claims made by all officers of the department.",
                "Claims for medical expenses and claims by external consultants are handled under separate procedures and are not covered here.",
            ],
        ),
    ]
    after = [
        (
            "5. Common Reasons for Rejection",
            [
                "Claims are most commonly returned because receipts are missing or illegible, the claim was submitted more than 30 days after the expense, or the amount exceeds the approved rate without justification.",
                "Returned claims may be corrected and resubmitted once. A claim returned twice must be discussed with the Finance Division before it is submitted again.",
            ],
        ),
        (
            "6. Records and Audit",
            [
                "Approved claims and their receipts are retained by the Finance Division for seven years. Officers should keep their own copies of receipts until payment is received.",
                "The Audit Unit samples at least five percent of approved claims every quarter. Officers must provide original receipts on request.",
            ],
        ),
    ]
    return before, after


def write_claims_pdf(path: Path, footer_text: str) -> None:
    styles = pdf_styles()
    title = "SOP: Claims Submission"
    story = header_block(title, ["Document No: SOP/FIN/2025/02", "Owner: Finance Division", "Version 1.0, effective 2 January 2025"], styles)
    before, after = claims_sections()
    story += sections_story(before, styles)
    story.append(Paragraph("3. Procedure", styles["h1"]))
    for number, step in enumerate(CLAIMS_STEPS, start=1):
        story.append(Paragraph(f"Step {number}. {step}", styles["body"]))
    story.append(Paragraph("4. Rates and Supporting Documents", styles["h1"]))
    story.append(Paragraph("The following rates apply to claims submitted under this SOP. Amounts above these rates require written justification from the Head of Division.", styles["body"]))
    story.append(styled_table(CLAIMS_RATES, [38 * mm, 36 * mm, 52 * mm, 38 * mm]))
    story.append(Spacer(1, 4 * mm))
    story += sections_story(after, styles)
    build_pdf(path, title, story, footer_text)


def write_circular_2024(path: Path) -> None:
    styles = pdf_styles()
    title = "Circular: Remote Work Arrangement"
    story = header_block(
        title,
        ["Reference: PKP/HR/2024/07", "Date: 2024-07-01", "Issued by: Human Resources Division"],
        styles,
    )
    story += sections_story(REMOTE_WORK_2024_SECTIONS, styles)
    story.append(Spacer(1, 6 * mm))
    story.append(Paragraph("Signed, Director, Jabatan Perkhidmatan Digital Negeri, 1 July 2024", styles["small"]))
    build_pdf(path, title, story, "PKP/HR/2024/07 | JPDN Human Resources Division")


def write_docx(path: Path, title: str, intro_lines: list[str], sections: list[Section]) -> None:
    document = DocxDocument()
    document.core_properties.title = title
    document.core_properties.author = AGENCY
    document.add_paragraph(AGENCY)
    document.add_heading(title, level=0)
    for line in intro_lines:
        document.add_paragraph(line)
    for heading, paragraphs in sections:
        document.add_heading(heading, level=1)
        for paragraph in paragraphs:
            document.add_paragraph(paragraph.replace("<b>", "").replace("</b>", ""))
    document.save(str(path))


def write_circular_2026(path: Path) -> None:
    write_docx(
        path,
        "Circular: Revised Remote Work Arrangement",
        [
            "Reference: PKP/HR/2026/03",
            "Effective date: 2026-03-01",
            "Issued by: Human Resources Division",
            "This circular supersedes PKP/HR/2024/07.",
        ],
        REMOTE_WORK_2026_SECTIONS,
    )


def write_procurement(path: Path, revised: bool) -> None:
    section_4 = PROCUREMENT_SECTION_4_V2 if revised else PROCUREMENT_SECTION_4_V1
    sections = PROCUREMENT_SECTIONS_COMMON + [section_4] + PROCUREMENT_SECTIONS_TAIL
    write_docx(
        path,
        "Procurement Policy",
        ["Reference: JPDN/FIN/POL/2025/01", "Effective date: 2025-01-15", "Owner: Finance Division"],
        sections,
    )


BUDGET_ROWS = [
    ["Cost Centre", "Category", "Approved Budget (RM)", "Q3 Allocation (RM)", "Actual Spend (RM)", "Variance (RM)", "Remarks"],
    ["FIN-100", "Staff salaries", 4200000, 1050000, 1032500, 17500, "Two vacancies unfilled"],
    ["FIN-110", "Overtime", 180000, 45000, 51200, -6200, "Year-end migration workload"],
    ["FIN-200", "Cloud hosting", 960000, 240000, 236400, 3600, "Within reserved capacity"],
    ["FIN-210", "Software licences", 520000, 130000, 141800, -11800, "Licence true-up for analytics suite"],
    ["FIN-300", "Training", 150000, 37500, 22300, 15200, "Two courses postponed to Q4"],
    ["FIN-310", "Travel and claims", 90000, 22500, 19850, 2650, "Remote work reduced travel"],
    ["FIN-400", "Procurement of hardware", 600000, 150000, 162000, -12000, "Laptop refresh brought forward"],
    ["FIN-500", "Professional services", 400000, 100000, 87500, 12500, "Audit fees billed in October"],
    ["FIN-600", "Connectivity allowance", 72000, 18000, 17650, 350, "RM50 monthly allowance"],
    ["FIN-900", "Contingency", 200000, 50000, 0, 50000, "Unused"],
]


def write_budget(path: Path) -> None:
    workbook = openpyxl.Workbook()
    workbook.properties.title = "Quarterly Budget Q3 2026"
    workbook.properties.creator = AGENCY
    sheet = workbook.active
    sheet.title = "Q3 2026 Summary"
    sheet["A1"] = "Quarterly Budget Q3 2026"
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = f"{AGENCY} - Finance Division - July to September 2026"
    for row_offset, row in enumerate(BUDGET_ROWS):
        for column_offset, value in enumerate(row):
            cell = sheet.cell(row=4 + row_offset, column=1 + column_offset, value=value)
            if row_offset == 0:
                cell.font = Font(bold=True)
    total_row = 4 + len(BUDGET_ROWS)
    sheet.cell(row=total_row, column=1, value="TOTAL")
    for column_index in range(3, 7):
        sheet.cell(row=total_row, column=column_index, value=sum(row[column_index - 1] for row in BUDGET_ROWS[1:]))
    notes = workbook.create_sheet("Notes")
    notes["A1"] = "Budget Notes Q3 2026"
    notes["A2"] = "Item"
    notes["B2"] = "Note"
    notes_rows = [
        ("Overtime", "Overtime exceeded allocation due to the claims system migration in August 2026."),
        ("Software licences", "True-up payment for the analytics platform approved by the Procurement Committee."),
        ("Contingency", "Contingency reserve remains untouched and will be reviewed in Q4."),
        ("Forecast", "Full-year spend is forecast at 97 percent of the approved budget."),
    ]
    for index, (item, note) in enumerate(notes_rows, start=3):
        notes.cell(row=index, column=1, value=item)
        notes.cell(row=index, column=2, value=note)
    workbook.save(str(path))


MINUTES_SECTIONS: list[Section] = [
    (
        "1. Kehadiran",
        [
            "Mesyuarat dipengerusikan oleh Pengarah Jabatan Perkhidmatan Digital Negeri. Turut hadir ialah Ketua Bahagian Kewangan, Ketua Bahagian Sumber Manusia, Ketua Unit Keselamatan ICT, Ketua Unit Integriti dan wakil Unit Perkhidmatan Digital.",
            "Seramai dua belas orang pegawai hadir. Ketua Unit Audit Dalam tidak dapat hadir kerana bertugas di luar daerah dan telah memohon maaf.",
        ],
    ),
    (
        "2. Pengesahan Minit Mesyuarat Lalu",
        [
            "Minit Mesyuarat Pengurusan Bil. 5/2026 yang diadakan pada 20 Mei 2026 telah dibentangkan dan disahkan tanpa pindaan, dicadangkan oleh Ketua Bahagian Kewangan dan disokong oleh Ketua Bahagian Sumber Manusia.",
        ],
    ),
    (
        "3. Perkara Berbangkit",
        [
            "Bahagian Sumber Manusia memaklumkan bahawa pelaksanaan bekerja dari rumah sebanyak tiga hari seminggu di bawah pekeliling PKP/HR/2026/03 berjalan lancar. Kadar kehadiran pada hari sauh adalah 94 peratus.",
            "Bahagian Kewangan memaklumkan bahawa elaun ketersambungan RM50 sebulan telah mula dibayar melalui tuntutan bulanan.",
        ],
    ),
    (
        "4. Laporan Kemajuan Projek Sistem Pengurusan Dokumen",
        [
            "Unit Perkhidmatan Digital membentangkan kemajuan projek sistem pengurusan dokumen pintar. Sistem ini membolehkan pegawai mencari dokumen dasar, pekeliling dan prosedur operasi standard dengan pantas menggunakan carian pintar.",
            "Mesyuarat bersetuju supaya fasa rintis dilaksanakan di Bahagian Kewangan dan Bahagian Sumber Manusia mulai Julai 2026. Akses kepada dokumen akan dikawal mengikut bahagian.",
        ],
    ),
    (
        "5. Prestasi Perbelanjaan Suku Kedua",
        [
            "Ketua Bahagian Kewangan melaporkan bahawa perbelanjaan sehingga akhir Jun 2026 adalah 48 peratus daripada peruntukan tahunan. Perbelanjaan lebih masa melebihi unjuran disebabkan migrasi sistem tuntutan.",
            "Mesyuarat mengarahkan semua bahagian mengemukakan unjuran perbelanjaan suku ketiga sebelum 15 Julai 2026.",
        ],
    ),
    (
        "6. Hal-hal Lain",
        [
            "Unit Keselamatan ICT mengingatkan semua pegawai supaya menggunakan VPN jabatan apabila bekerja dari rumah dan melaporkan sebarang insiden keselamatan dalam tempoh dua jam.",
            "Mesyuarat ditangguhkan pada jam 12.30 tengah hari. Mesyuarat seterusnya dijadualkan pada 22 Julai 2026.",
        ],
    ),
]


def write_minutes(path: Path) -> None:
    styles = pdf_styles()
    title = "Minit Mesyuarat Pengurusan Bil. 6/2026"
    story = header_block(
        title,
        ["Tarikh: 24 Jun 2026", "Masa: 9.30 pagi", "Tempat: Bilik Mesyuarat Utama, Aras 5, JPDN"],
        styles,
    )
    story += sections_story(MINUTES_SECTIONS, styles)
    story.append(Spacer(1, 6 * mm))
    story.append(Paragraph("Disediakan oleh: Setiausaha Mesyuarat. Disahkan oleh: Pengarah JPDN.", styles["small"]))
    build_pdf(path, title, story, "Minit Mesyuarat Pengurusan Bil. 6/2026 | SULIT")


ZH_LINES = [
    ("数据共享指南", 54),
    ("数字服务部 (JPDN) 公共指南 编号 JPDN/DG/2026/01", 24),
    ("", 20),
    ("一、目的", 32),
    ("本指南规定各部门之间共享数据的原则和程序，", 26),
    ("确保个人数据受到保护，并提高公共服务效率。", 26),
    ("二、数据分类", 32),
    ("所有数据分为公开、内部、受限和机密四个等级。", 26),
    ("受限及机密数据不得通过电子邮件发送。", 26),
    ("三、共享程序", 32),
    ("1. 申请部门提交数据共享申请表。", 26),
    ("2. 数据拥有者在五个工作日内审核申请。", 26),
    ("3. 数据治理委员会批准后签署数据共享协议。", 26),
    ("4. 数据必须通过加密渠道传输。", 26),
    ("四、个人数据保护", 32),
    ("共享个人数据前必须进行匿名化处理，", 26),
    ("并遵守个人数据保护法的规定。", 26),
    ("五、生效日期", 32),
    ("本指南自2026年4月1日起生效。", 26),
]


def write_chinese_png(path: Path) -> None:
    width, height = 1240, 1754
    image = Image.new("RGB", (width, height), (250, 248, 242))
    draw = ImageDraw.Draw(image)
    y = 110
    for text, size in ZH_LINES:
        if text:
            font = load_font(CJK_FONT_CANDIDATES, size)
            draw.text((100, y), text, font=font, fill=(25, 25, 30))
        y += int(size * 1.9)
    image = image.rotate(0.6, resample=Image.Resampling.BICUBIC, expand=False, fillcolor=(250, 248, 242))
    image = image.filter(ImageFilter.GaussianBlur(radius=0.6))
    image.save(str(path), format="PNG")


REPORT_PAGES = [
    (
        "Annual Report on Digital Services 2025",
        [
            "Executive Summary",
            "In 2025 Jabatan Perkhidmatan Digital Negeri delivered 42 online services to the public, up from 31 in 2024. Total online transactions reached 1.8 million, and the average processing time for licence applications fell from nine days to four days.",
            "The department completed the migration of its core systems to the state government cloud, launched a single sign-on service for citizens, and began a pilot of an intelligent document management system for internal policies and circulars.",
        ],
    ),
    (
        "Service Adoption and Customer Satisfaction",
        [
            "The online business licence renewal service was the most used service, with 412,000 transactions in 2025.",
            "Customer satisfaction across all digital services averaged 4.3 out of 5 based on 26,000 survey responses.",
            "Mobile access accounted for 61 percent of all sessions, compared with 48 percent in 2024.",
            "The most common complaints were slow payment confirmation and unclear error messages.",
            "Average helpdesk response time improved from 14 hours to 6 hours after the chatbot launch.",
        ],
    ),
    (
        "Infrastructure, Security and Cost",
        [
            "Service availability across the year was 99.93 percent, exceeding the target of 99.5 percent. Two planned maintenance windows and one unplanned outage of 47 minutes were recorded.",
            "The ICT Security Unit handled 38 security incidents, none of which resulted in the loss of personal data. All officers completed the annual phishing awareness exercise.",
            "Operating cost per online transaction fell to RM1.20 from RM1.85 in 2024, mainly due to cloud consolidation and reduced licence costs.",
        ],
    ),
    (
        "Priorities for 2026",
        [
            "Expand the digital document management system to all divisions with department-based access control.",
            "Introduce Malay, English and Chinese search across policies, circulars and meeting minutes.",
            "Publish an open data catalogue in line with the data sharing guideline.",
            "Reduce average licence processing time to two days.",
            "Train 300 officers in data governance and information security.",
        ],
    ),
]


def rasterized_page_image(heading: str, lines: list[str]) -> Image.Image:
    width, height = 1654, 2339
    image = Image.new("L", (width, height), 246)
    draw = ImageDraw.Draw(image)
    heading_font = load_font(LATIN_FONT_CANDIDATES, 52)
    body_font = load_font(LATIN_FONT_CANDIDATES, 34)
    draw.text((140, 160), heading, font=heading_font, fill=20)
    y = 300
    for line in lines:
        words = line.split()
        current = ""
        for word in words:
            candidate = f"{current} {word}".strip()
            if draw.textlength(candidate, font=body_font) > width - 280:
                draw.text((140, y), current, font=body_font, fill=30)
                y += 50
                current = word
            else:
                current = candidate
        if current:
            draw.text((140, y), current, font=body_font, fill=30)
            y += 50
        y += 30
    image = image.rotate(0.4, resample=Image.Resampling.BICUBIC, fillcolor=246)
    return image.filter(ImageFilter.GaussianBlur(radius=0.5))


def write_scanned_report(path: Path) -> None:
    page_width, page_height = A4
    canvas = pdf_canvas.Canvas(str(path), pagesize=A4, invariant=1)
    canvas.setTitle("Annual Report on Digital Services 2025")
    canvas.setAuthor(AGENCY)
    for page_number, (heading, lines) in enumerate(REPORT_PAGES, start=1):
        if page_number in (2, 4):
            image = rasterized_page_image(heading, lines)
            canvas.drawImage(ImageReader(image), 0, 0, width=page_width, height=page_height)
        else:
            canvas.setFont("Helvetica", 9)
            canvas.drawString(20 * mm, page_height - 15 * mm, AGENCY)
            canvas.setFont("Helvetica-Bold", 18 if page_number == 1 else 15)
            canvas.drawString(20 * mm, page_height - 30 * mm, heading)
            text = canvas.beginText(20 * mm, page_height - 45 * mm)
            text.setFont("Helvetica", 11)
            text.setLeading(15)
            for line in lines:
                words = line.split()
                current = ""
                for word in words:
                    candidate = f"{current} {word}".strip()
                    if canvas.stringWidth(candidate, "Helvetica", 11) > page_width - 40 * mm:
                        text.textLine(current)
                        current = word
                    else:
                        current = candidate
                if current:
                    text.textLine(current)
                text.textLine("")
            canvas.drawText(text)
            canvas.setFont("Helvetica", 8)
            canvas.drawRightString(page_width - 20 * mm, 12 * mm, f"Page {page_number}")
        canvas.showPage()
    canvas.save()


MANIFEST = [
    {"file": "circular_2024_07_remote_work.pdf", "department_id": "hr", "uploader": "ben"},
    {"file": "circular_2026_03_remote_work.docx", "department_id": "hr", "uploader": "ben"},
    {"file": "policy_procurement.docx", "department_id": "finance", "uploader": "alice"},
    {"file": "sop_claims_submission.pdf", "department_id": "finance", "uploader": "alice"},
    {"file": "quarterly_budget_q3_2026.xlsx", "department_id": "finance", "uploader": "alice"},
    {"file": "minutes_mesyuarat_jun_2026.pdf", "department_id": "public", "uploader": "admin"},
    {"file": "guideline_data_sharing_zh.png", "department_id": "public", "uploader": "admin"},
    {"file": "report_digital_services_2025_scanned.pdf", "department_id": "public", "uploader": "admin"},
    {"file": "sop_claims_submission_copy.pdf", "department_id": "finance", "uploader": "alice"},
    {"file": "sop_claims_submission_resaved.pdf", "department_id": "finance", "uploader": "alice"},
    {"file": "policy_procurement_v2.docx", "department_id": "finance", "uploader": "alice", "version_of": "policy_procurement.docx"},
]


def main() -> None:
    CORPUS_DIR.mkdir(parents=True, exist_ok=True)
    write_circular_2024(CORPUS_DIR / "circular_2024_07_remote_work.pdf")
    write_circular_2026(CORPUS_DIR / "circular_2026_03_remote_work.docx")
    write_procurement(CORPUS_DIR / "policy_procurement.docx", revised=False)
    write_procurement(CORPUS_DIR / "policy_procurement_v2.docx", revised=True)
    write_claims_pdf(CORPUS_DIR / "sop_claims_submission.pdf", "SOP/FIN/2025/02 | Finance Division | Printed copy is uncontrolled")
    shutil.copyfile(CORPUS_DIR / "sop_claims_submission.pdf", CORPUS_DIR / "sop_claims_submission_copy.pdf")
    write_claims_pdf(CORPUS_DIR / "sop_claims_submission_resaved.pdf", "SOP/FIN/2025/02 | Finance Division | Re-issued copy, reviewed March 2026")
    write_budget(CORPUS_DIR / "quarterly_budget_q3_2026.xlsx")
    write_minutes(CORPUS_DIR / "minutes_mesyuarat_jun_2026.pdf")
    write_chinese_png(CORPUS_DIR / "guideline_data_sharing_zh.png")
    write_scanned_report(CORPUS_DIR / "report_digital_services_2025_scanned.pdf")
    (CORPUS_DIR / "manifest.json").write_text(json.dumps(MANIFEST, indent=2) + "\n", encoding="utf-8")
    for entry in MANIFEST:
        target = CORPUS_DIR / entry["file"]
        print(f"{target.stat().st_size:>9}  {entry['file']}")


if __name__ == "__main__":
    main()
